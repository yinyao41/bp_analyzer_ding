import streamlit as st
import pandas as pd
from pypdf import PdfReader
from pptx import Presentation
import dashscope
import os
from docx import Document

# ==========================
# 获取 API Key
# ==========================
dashscope.api_key = os.environ.get("DASHSCOPE_API_KEY")

# ==========================
# 页面设置（已去除红框提示）
# ==========================
st.set_page_config(page_title="科技项目初筛分析", layout="wide")
st.title("🔍 科技项目初筛分析")

# ==========================
# 加载内置知识库（BP提示词 + TBS-V2 + 调研要点）
# ==========================
@st.cache_data
def load_builtin_knowledge():
    try:
        # BP提示词模板（核心输出格式）
        bp_doc = Document("BP提示词内容.docx")
        bp_template = "\n".join(p.text.strip() for p in bp_doc.paragraphs if p.text.strip())
        
        # TBS-V2
        tbs = pd.read_excel("TBS-V2.xlsx", sheet_name="Sheet2")
        
        # 调研要点
        survey_df = pd.read_excel("调研要点.xlsx", sheet_name=0)
        survey_text = survey_df.to_string(index=False)
        
        return {
            "bp_template": bp_template,
            "tbs_text": tbs.to_string(index=False),
            "survey_text": survey_text
        }
    except Exception as e:
        st.error(f"知识库加载失败: {e}")
        st.stop()

BUILTIN_KNOWLEDGE = load_builtin_knowledge()

# ==========================
# 文件读取函数
# ==========================
def extract_text_from_file(uploaded_file):
    if not uploaded_file:
        return ""
    ext = os.path.splitext(uploaded_file.name)[1].lower()
    bytes_data = uploaded_file.read()
    uploaded_file.seek(0)
    
    if ext == ".pdf":
        reader = PdfReader(bytes_data)
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    elif ext == ".pptx":
        prs = Presentation(bytes_data)
        return "\n".join(shape.text for slide in prs.slides for shape in slide.shapes if hasattr(shape, "text"))
    return ""

# ==========================
# 上传区域（简洁版）
# ==========================
uploaded_file = st.file_uploader(
    "上传项目资料（PDF / PPTX）",
    type=["pdf", "pptx"],
    help="建议文件不超过20MB"
)

# ==========================
# 分析流程
# ==========================
if uploaded_file:
    with st.spinner("正在提取文档内容..."):
        doc_text = extract_text_from_file(uploaded_file)
    
    st.success("✅ 文档内容提取完成")
    
    if st.button("🚀 开始初筛分析", type="primary"):
        prompt = f"""{BUILTIN_KNOWLEDGE['bp_template']}

用户上传的BP文档内容（请严格基于此内容进行分析，不得编造任何未提及信息）：
{doc_text[:15000]}

内置 TBS-V2 规则摘要：
{BUILTIN_KNOWLEDGE['tbs_text'][:10000]}

内置调研要点摘要：
{BUILTIN_KNOWLEDGE['survey_text'][:5000]}
"""

        with st.spinner("AI 正在生成结构化初筛报告..."):
            try:
                response = dashscope.Generation.call(
                    model="qwen-max",
                    prompt=prompt,
                    temperature=0.1,
                    max_tokens=12000
                )
                result = response.output.text
            except Exception as e:
                st.error(f"分析失败: {e}")
                result = ""

        if result:
            st.subheader("📋 项目初筛分析报告")
            st.markdown(result)
            
            # 下载按钮
            st.download_button(
                label="📥 下载完整报告（Markdown）",
                data=result,
                file_name="项目初筛分析报告.md",
                mime="text/markdown"
            )
else:
    st.info("请上传项目资料文件（PDF 或 PPTX）")
